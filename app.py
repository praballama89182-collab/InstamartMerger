import os
import io
import pandas as pd
import numpy as np
import streamlit as st
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

st.set_page_config(
    page_title="Swiggy Granular Ad Analytics Dashboard",
    page_icon="🔵",
    layout="wide"
)

st.markdown("""
<style>
    :root {
        --primary-blue: #1E88E5;
        --primary-blue-dark: #1565C0;
        --primary-blue-light: #E3F2FD;
        --primary-border: #BBDEFB;
        --primary-text: #3D3D3D;
    }

    .stApp {
        background: linear-gradient(180deg, #F4F8FB 0%, #FFFFFF 40%);
        color: var(--primary-text);
    }

    [data-testid="stHeader"] {
        background: rgba(255,255,255,0.92);
    }

    h1, h2, h3 {
        color: var(--primary-blue-dark) !important;
        font-weight: 700 !important;
    }

    [data-testid="stMetric"] {
        background: linear-gradient(135deg, #FFFFFF 0%, #E3F2FD 100%);
        border: 1px solid var(--primary-border);
        border-left: 5px solid var(--primary-blue);
        border-radius: 12px;
        padding: 14px 18px;
        box-shadow: 0 3px 12px rgba(30, 136, 229, 0.10);
    }

    [data-testid="stMetricLabel"] {
        color: #5A6B7C !important;
        font-weight: 600 !important;
    }

    [data-testid="stMetricValue"] {
        color: var(--primary-blue-dark) !important;
        font-weight: 750 !important;
    }

    .stTabs [data-baseweb="tab-list"] {
        gap: 5px;
        background: #E3F2FD;
        border: 1px solid var(--primary-border);
        border-radius: 12px;
        padding: 5px;
    }

    .stTabs [data-baseweb="tab"] {
        height: 42px;
        border-radius: 7px;
        color: #1565C0;
        font-weight: 600;
    }

    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #1E88E5 0%, #1565C0 100%) !important;
        color: #FFFFFF !important;
        box-shadow: 0 2px 7px rgba(30, 136, 229, 0.25);
    }

    [data-testid="stFileUploader"] {
        background: #FFFFFF;
        border: 1px solid var(--primary-border);
        border-radius: 12px;
        padding: 8px;
        box-shadow: 0 2px 8px rgba(30, 136, 229, 0.07);
    }

    [data-testid="stFileUploaderDropzone"] {
        background: #F9FBFC;
        border: 1px dashed #90CAF9;
        border-radius: 9px;
    }

    [data-testid="stFileUploader"] button,
    .stDownloadButton button,
    .stButton button {
        background: linear-gradient(135deg, #1E88E5, #1565C0) !important;
        color: #FFFFFF !important;
        border: 0 !important;
        border-radius: 8px !important;
        font-weight: 650 !important;
        box-shadow: 0 2px 7px rgba(30, 136, 229, 0.20);
    }

    [data-testid="stFileUploader"] button:hover,
    .stDownloadButton button:hover,
    .stButton button:hover {
        background: #0D47A1 !important;
        color: #FFFFFF !important;
    }

    [data-testid="stDataFrame"] {
        border: 1px solid var(--primary-border);
        border-radius: 10px;
        overflow: hidden;
        box-shadow: 0 2px 9px rgba(30, 136, 229, 0.06);
    }

    [data-testid="stDataFrame"] th,
    [data-testid="stDataFrame"] td {
        text-align: center !important;
        vertical-align: middle !important;
    }

    [data-testid="stDataFrame"] th {
        background: #BBDEFB !important;
        color: #0D47A1 !important;
        font-weight: 700 !important;
    }

    hr {
        border-color: #D0E1F9 !important;
    }
</style>
""", unsafe_allow_html=True)

st.title("🔵 Swiggy Granular Ad Analytics & Merger Dashboard")
st.write("Upload Swiggy granular advertising spreadsheets or CSV files. Inspect raw data, consolidated master tables, Month-on-Month comparison sheets, trend analytics, city performance, product performance, and ad property insights.")

def style_and_export_pivot(pivot_df, sheet_name="Comparison"):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    
    top_header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    sec_header_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="1565C0")
    
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    
    data_font = Font(name="Calibri", size=10, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    if isinstance(pivot_df.columns, pd.MultiIndex):
        metrics = pivot_df.columns.get_level_values(1)
        entity_title = pivot_df.index.name if pivot_df.index.name else ""

        ws.cell(row=1, column=1, value="").fill = top_header_fill
        ws.cell(row=1, column=1).border = thin_border
        
        current_col = 2
        for month in pivot_df.columns.get_level_values(0).unique():
            sub_cols = [c for c in pivot_df.columns if c[0] == month]
            num_sub = len(sub_cols)
            if num_sub > 1:
                ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
            cell = ws.cell(row=1, column=current_col, value=str(month))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            
            for c in range(current_col, current_col + num_sub):
                ws.cell(row=1, column=c).border = thin_border
                ws.cell(row=1, column=c).fill = top_header_fill
            current_col += num_sub

        cell_a2 = ws.cell(row=2, column=1, value=str(entity_title))
        cell_a2.fill = index_fill
        cell_a2.font = index_font
        cell_a2.alignment = align_center
        cell_a2.border = thin_border

        for col_idx, metric in enumerate(metrics, start=2):
            cell = ws.cell(row=2, column=col_idx, value=str(metric))
            cell.fill = sec_header_fill
            cell.font = sec_header_font
            cell.alignment = align_center
            cell.border = thin_border

        start_data_row = 3
        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=start_data_row):
            is_grand_total = (str(idx_val).strip().lower() == 'grand total')
            
            idx_cell = ws.cell(row=r_idx, column=1, value=str(idx_val))
            idx_cell.fill = total_fill if is_grand_total else index_fill
            idx_cell.font = total_font if is_grand_total else index_font
            idx_cell.alignment = align_left
            idx_cell.border = thin_border

            for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_grand_total else data_font
                if is_grand_total:
                    val_cell.fill = total_fill
                val_cell.alignment = align_center
                val_cell.border = thin_border

                if isinstance(val, (int, float, np.number)):
                    if str(m_col).upper() == 'ACOS':
                        val_cell.value = val / 100.0 if val > 1 else val
                        val_cell.number_format = '0.00%'
                    elif str(m_col).upper() == 'CPM':
                        val_cell.value = round(val)
                        val_cell.number_format = '#,##0'
                    else:
                        val_cell.value = round(val, 2)
                        val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                else:
                    val_cell.value = val
    else:
        for col_idx, col_name in enumerate(pivot_df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=str(col_name))
            cell.fill = top_header_fill
            cell.font = top_header_font
            cell.alignment = align_center
            cell.border = thin_border

        for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
            is_pct_row = ('VS' in str(row_data.iloc[0]).strip().upper())
            
            for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                val_cell = ws.cell(row=r_idx, column=c_idx)
                val_cell.font = total_font if is_pct_row else data_font
                val_cell.alignment = align_center
                val_cell.border = thin_border
                
                if str(col_name).upper() == 'ACOS' and not is_pct_row and isinstance(val, (int, float, np.number)):
                    val_cell.value = val / 100.0 if val > 1 else val
                    val_cell.number_format = '0.00%'
                elif str(col_name).upper() == 'CPM' and not is_pct_row and isinstance(val, (int, float, np.number)):
                    val_cell.value = round(val)
                    val_cell.number_format = '#,##0'
                elif isinstance(val, float):
                    val_cell.value = round(val, 2)
                    val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                elif isinstance(val, (int, np.integer)):
                    val_cell.value = val
                    val_cell.number_format = '#,##0'
                else:
                    val_cell.value = val

                if is_pct_row and c_idx > 1:
                    val_str = str(val).replace('%', '').replace('+', '').strip()
                    try:
                        num_v = float(val_str)
                        if num_v < 0:
                            val_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                        elif num_v > 0:
                            val_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                            val_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                    except ValueError:
                        pass
                elif is_pct_row:
                    val_cell.fill = sec_header_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def convert_all_pivots_to_excel(pivot_dict):
    buffer = io.BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    top_header_fill = PatternFill(start_color="1E88E5", end_color="1E88E5", fill_type="solid")
    top_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    sec_header_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    sec_header_font = Font(name="Calibri", size=10, bold=True, color="1565C0")
    index_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    index_font = Font(name="Calibri", size=10, bold=True, color="000000")
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=False)
    thin_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF')
    )

    for sheet_name, pivot_df in pivot_dict.items():
        if pivot_df is not None and not pivot_df.empty:
            ws = wb.create_sheet(title=sheet_name[:31])

            if isinstance(pivot_df.columns, pd.MultiIndex):
                metrics = pivot_df.columns.get_level_values(1)
                entity_title = pivot_df.index.name if pivot_df.index.name else ""

                ws.cell(row=1, column=1, value="").fill = top_header_fill
                ws.cell(row=1, column=1).border = thin_border
                
                current_col = 2
                for month in pivot_df.columns.get_level_values(0).unique():
                    sub_cols = [c for c in pivot_df.columns if c[0] == month]
                    num_sub = len(sub_cols)
                    if num_sub > 1:
                        ws.merge_cells(start_row=1, start_column=current_col, end_row=1, end_column=current_col + num_sub - 1)
                    cell = ws.cell(row=1, column=current_col, value=str(month))
                    cell.fill = top_header_fill
                    cell.font = top_header_font
                    cell.alignment = align_center
                    
                    for c in range(current_col, current_col + num_sub):
                        ws.cell(row=1, column=c).border = thin_border
                        ws.cell(row=1, column=c).fill = top_header_fill
                    current_col += num_sub

                cell_a2 = ws.cell(row=2, column=1, value=str(entity_title))
                cell_a2.fill = index_fill
                cell_a2.font = index_font
                cell_a2.alignment = align_center
                cell_a2.border = thin_border

                for col_idx, metric in enumerate(metrics, start=2):
                    cell = ws.cell(row=2, column=col_idx, value=str(metric))
                    cell.fill = sec_header_fill
                    cell.font = sec_header_font
                    cell.alignment = align_center
                    cell.border = thin_border

                for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=3):
                    is_grand_total = (str(idx_val).strip().lower() == 'grand total')
                    
                    idx_cell = ws.cell(row=r_idx, column=1, value=str(idx_val))
                    idx_cell.fill = total_fill if is_grand_total else index_fill
                    idx_cell.font = total_font if is_grand_total else index_font
                    idx_cell.alignment = align_left
                    idx_cell.border = thin_border

                    for c_idx, (m_col, val) in enumerate(zip(metrics, row_data), start=2):
                        val_cell = ws.cell(row=r_idx, column=c_idx)
                        val_cell.font = total_font if is_grand_total else data_font
                        if is_grand_total:
                            val_cell.fill = total_fill
                        val_cell.alignment = align_center
                        val_cell.border = thin_border

                        if isinstance(val, (int, float, np.number)):
                            if str(m_col).upper() == 'ACOS':
                                val_cell.value = val / 100.0 if val > 1 else val
                                val_cell.number_format = '0.00%'
                            elif str(m_col).upper() == 'CPM':
                                val_cell.value = round(val)
                                val_cell.number_format = '#,##0'
                            else:
                                val_cell.value = round(val, 2)
                                val_cell.number_format = '#,##0' if float(val).is_integer() else '#,##0.00'
                        else:
                            val_cell.value = val
            else:
                for col_idx, col_name in enumerate(pivot_df.columns, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=str(col_name))
                    cell.fill = top_header_fill
                    cell.font = top_header_font
                    cell.alignment = align_center
                    cell.border = thin_border

                for r_idx, (idx_val, row_data) in enumerate(pivot_df.iterrows(), start=2):
                    is_pct_row = ('VS' in str(row_data.iloc[0]).strip().upper())
                    for c_idx, (col_name, val) in enumerate(zip(pivot_df.columns, row_data), start=1):
                        val_cell = ws.cell(row=r_idx, column=c_idx)
                        val_cell.font = total_font if is_pct_row else data_font
                        val_cell.alignment = align_center
                        val_cell.border = thin_border
                        
                        if str(col_name).upper() == 'ACOS' and not is_pct_row and isinstance(val, (int, float, np.number)):
                            val_cell.value = val / 100.0 if val > 1 else val
                            val_cell.number_format = '0.00%'
                        elif str(col_name).upper() == 'CPM' and not is_pct_row and isinstance(val, (int, float, np.number)):
                            val_cell.value = round(val)
                            val_cell.number_format = '#,##0'
                        elif isinstance(val, float):
                            val_cell.value = round(val, 2)
                            val_cell.number_format = '#,##0.00'
                        elif isinstance(val, (int, np.integer)):
                            val_cell.value = val
                            val_cell.number_format = '#,##0'
                        else:
                            val_cell.value = val

                        if is_pct_row and c_idx > 1:
                            val_str = str(val).replace('%', '').replace('+', '').strip()
                            try:
                                num_v = float(val_str)
                                if num_v < 0:
                                    val_cell.fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
                                    val_cell.font = Font(name="Calibri", size=10, bold=True, color="721C24")
                                elif num_v > 0:
                                    val_cell.fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
                                    val_cell.font = Font(name="Calibri", size=10, bold=True, color="155724")
                            except ValueError:
                                pass
                        elif is_pct_row:
                            val_cell.fill = sec_header_fill

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

col1, col2, col3, col4, col5 = st.columns(5)
uploaded_files = []

with col1:
    f1 = st.file_uploader("Upload File 1", type=["csv", "xlsx", "xls"], key="file1")
    if f1: uploaded_files.append(f1)

with col2:
    f2 = st.file_uploader("Upload File 2", type=["csv", "xlsx", "xls"], key="file2")
    if f2: uploaded_files.append(f2)

with col3:
    f3 = st.file_uploader("Upload File 3", type=["csv", "xlsx", "xls"], key="file3")
    if f3: uploaded_files.append(f3)

with col4:
    f4 = st.file_uploader("Upload File 4", type=["csv", "xlsx", "xls"], key="file4")
    if f4: uploaded_files.append(f4)

with col5:
    f5 = st.file_uploader("Upload File 5", type=["csv", "xlsx", "xls"], key="file5")
    if f5: uploaded_files.append(f5)

if uploaded_files:
    consolidated_dfs = []
    raw_files_dict = {}
    consolidated_raw_tabs = {}

    for uploaded_file in uploaded_files:
        fallback_month_name = os.path.splitext(uploaded_file.name)[0].upper()
        raw_files_dict[uploaded_file.name] = {}
        
        try:
            if uploaded_file.name.endswith('.csv'):
                df = pd.read_csv(uploaded_file, skiprows=6)
                sheet_name = "Sheet1"
                raw_files_dict[uploaded_file.name][sheet_name] = df.copy()
                
                date_col = 'METRICS_DATE' if 'METRICS_DATE' in df.columns else None
                if date_col:
                    dt_series = pd.to_datetime(df[date_col], errors='coerce')
                    month_series = dt_series.dt.strftime('%B').str.upper()
                    df['Month'] = month_series.fillna(fallback_month_name)
                else:
                    df['Month'] = fallback_month_name

                df_consolidated = df.copy()
                if 'Month' in df_consolidated.columns:
                    col_month = df_consolidated.pop('Month')
                    df_consolidated.insert(0, 'Month', col_month)
                else:
                    df_consolidated.insert(0, 'Month', fallback_month_name)
                    
                df_consolidated.insert(1, 'Tab Name', sheet_name)
                consolidated_dfs.append(df_consolidated)
                
                if sheet_name not in consolidated_raw_tabs:
                    consolidated_raw_tabs[sheet_name] = []
                consolidated_raw_tabs[sheet_name].append(df.copy())
            else:
                xls = pd.ExcelFile(uploaded_file)
                for sheet_name in xls.sheet_names:
                    df = pd.read_excel(xls, sheet_name=sheet_name)
                    raw_files_dict[uploaded_file.name][sheet_name] = df.copy()
                    
                    date_col = 'METRICS_DATE' if 'METRICS_DATE' in df.columns else ('Date' if 'Date' in df.columns else None)
                    if date_col:
                        dt_series = pd.to_datetime(df[date_col], errors='coerce')
                        month_series = dt_series.dt.strftime('%B').str.upper()
                        df['Month'] = month_series.fillna(fallback_month_name)
                    else:
                        df['Month'] = fallback_month_name

                    if sheet_name not in consolidated_raw_tabs:
                        consolidated_raw_tabs[sheet_name] = []
                    consolidated_raw_tabs[sheet_name].append(df.copy())
                    
                    df_consolidated = df.copy()
                    if 'Month' in df_consolidated.columns:
                        col_month = df_consolidated.pop('Month')
                        df_consolidated.insert(0, 'Month', col_month)
                    else:
                        df_consolidated.insert(0, 'Month', fallback_month_name)
                        
                    df_consolidated.insert(1, 'Tab Name', sheet_name)
                    consolidated_dfs.append(df_consolidated)
        except Exception as e:
            st.error(f"Error processing {uploaded_file.name}: {str(e)}")

    if consolidated_dfs:
        final_df = pd.concat(consolidated_dfs, ignore_index=True)
        
        def get_numeric_col(df, possible_cols):
            for col in possible_cols:
                if col in df.columns:
                    return pd.to_numeric(df[col], errors='coerce').fillna(0)
            return pd.Series(0, index=df.index)

        final_df['_impressions'] = get_numeric_col(final_df, ['TOTAL_IMPRESSIONS', 'Impressions'])
        final_df['_atc'] = get_numeric_col(final_df, ['TOTAL_A2C', 'A2C'])
        final_df['_orders'] = get_numeric_col(final_df, ['TOTAL_CONVERSIONS', 'Orders'])
        final_df['_sales'] = get_numeric_col(final_df, ['TOTAL_GMV', 'GMV', 'Sales'])
        final_df['_budget_consumed'] = get_numeric_col(final_df, ['TOTAL_BUDGET_BURNT', 'Budget Burnt', 'Spend'])

        final_df['ROAS'] = final_df.apply(lambda r: round(r['_sales'] / r['_budget_consumed'], 2) if r['_budget_consumed'] > 0 else 0.0, axis=1)
        final_df['ACOS'] = final_df.apply(lambda r: round((r['_budget_consumed'] / r['_sales']) * 100, 2) if r['_sales'] > 0 else 0.0, axis=1)
        final_df['CPM'] = final_df.apply(lambda r: round((r['_budget_consumed'] / r['_impressions']) * 1000) if r['_impressions'] > 0 else 0, axis=1)

        date_col = 'METRICS_DATE' if 'METRICS_DATE' in final_df.columns else ('Date' if 'Date' in final_df.columns else None)
        if date_col:
            final_df['_date_dt'] = pd.to_datetime(final_df[date_col], errors='coerce')
            def assign_week_formatted(row):
                dt = row['_date_dt']
                if pd.isna(dt): return np.nan
                day = dt.day
                if 1 <= day <= 7: return "Week 1 (1-7)"
                elif 8 <= day <= 14: return "Week 2 (8-14)"
                elif 15 <= day <= 21: return "Week 3 (15-21)"
                elif 22 <= day <= 28: return "Week 4 (22-28)"
                elif day >= 29: return "Week 5 (29-31)"
                return np.nan
            final_df['Week'] = final_df.apply(assign_week_formatted, axis=1)
        else:
            final_df['Week'] = np.nan

        MONTH_ORDER = {
            "JANUARY": 1, "FEBRUARY": 2, "MARCH": 3, "APRIL": 4,
            "MAY": 5, "JUNE": 6, "JULY": 7, "AUGUST": 8,
            "SEPTEMBER": 9, "OCTOBER": 10, "NOVEMBER": 11, "DECEMBER": 12
        }

        def get_calendar_months(df_input):
            months = [str(x).strip().upper() for x in df_input['Month'].dropna().unique()]
            return sorted(months, key=lambda x: MONTH_ORDER.get(x, 99))

        def compute_grouped_table(df_subset, group_col, selected_item="All"):
            if group_col not in df_subset.columns:
                return pd.DataFrame()
            
            df_working = df_subset.dropna(subset=[group_col]).copy()
            df_working[group_col] = df_working[group_col].astype(str)
            
            if selected_item and selected_item != "All":
                df_working = df_working[df_working[group_col] == selected_item]
            
            if df_working.empty:
                return pd.DataFrame()

            grouped = df_working.groupby(group_col).agg(
                IMPRESSIONS=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                ORDERS=('_orders', 'sum'),
                SPENDS=('_budget_consumed', 'sum'),
                SALES=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['SPENDS'] / r['IMPRESSIONS']) * 1000) if r['IMPRESSIONS'] > 0 else 0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['SALES'] / r['SPENDS'], 2) if r['SPENDS'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['SPENDS'] / r['SALES']) * 100, 2) if r['SALES'] > 0 else 0.0, axis=1)

            display_name = group_col.upper()
            grouped = grouped.rename(columns={group_col: display_name})
            col_order = [display_name, 'IMPRESSIONS', 'CPM', 'ATC', 'ORDERS', 'SPENDS', 'SALES', 'ROAS', 'ACOS']
            
            res_df = grouped.reindex(columns=col_order)
            res_df['SPENDS'] = res_df['SPENDS'].round(2)
            res_df['SALES'] = res_df['SALES'].round(2)
            return res_df

        def create_mom_comparison_table(df_input, entity_col):
            if entity_col not in df_input.columns:
                return pd.DataFrame()
            
            working_df = df_input.dropna(subset=[entity_col, 'Month']).copy()
            if working_df.empty:
                return pd.DataFrame()
            
            grouped = working_df.groupby([entity_col, 'Month']).agg(
                Impressions=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                Orders=('_orders', 'sum'),
                Spends=('_budget_consumed', 'sum'),
                Sales=('_sales', 'sum')
            ).reset_index()

            grouped['CPM'] = grouped.apply(lambda r: round((r['Spends'] / r['Impressions']) * 1000) if r['Impressions'] > 0 else 0, axis=1)
            grouped['ROAS'] = grouped.apply(lambda r: round(r['Sales'] / r['Spends'], 2) if r['Spends'] > 0 else 0.0, axis=1)
            grouped['ACOS'] = grouped.apply(lambda r: round((r['Spends'] / r['Sales']) * 100, 2) if r['Sales'] > 0 else 0.0, axis=1)

            pivot_df = grouped.pivot(
                index=entity_col,
                columns='Month',
                values=['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            )

            metrics_order = ['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            all_months = get_calendar_months(df_input)
            
            pivot_df = pivot_df.reorder_levels([1, 0], axis=1)
            sorted_cols = pd.MultiIndex.from_product([all_months, metrics_order], names=['Month', 'Metric'])
            pivot_df = pivot_df.reindex(columns=sorted_cols).fillna(0)

            grand_total_series = {}
            for month in all_months:
                month_df = working_df[working_df['Month'] == month]
                total_imp = month_df['_impressions'].sum()
                total_atc = month_df['_atc'].sum()
                total_orders = month_df['_orders'].sum()
                total_spends = round(month_df['_budget_consumed'].sum(), 2)
                total_sales = round(month_df['_sales'].sum(), 2)

                total_cpm = round((total_spends / total_imp) * 1000) if total_imp > 0 else 0
                total_roas = round(total_sales / total_spends, 2) if total_spends > 0 else 0.0
                total_acos = round((total_spends / total_sales) * 100, 2) if total_sales > 0 else 0.0

                grand_total_series[(month, 'Impressions')] = total_imp
                grand_total_series[(month, 'CPM')] = total_cpm
                grand_total_series[(month, 'ATC')] = total_atc
                grand_total_series[(month, 'Orders')] = total_orders
                grand_total_series[(month, 'Spends')] = total_spends
                grand_total_series[(month, 'Sales')] = total_sales
                grand_total_series[(month, 'ROAS')] = total_roas
                grand_total_series[(month, 'ACOS')] = total_acos

            pivot_df.loc['Grand Total'] = grand_total_series
            pivot_df.index.name = entity_col
            return pivot_df

        def create_monthly_summary_table(df_input):
            working_df = df_input.dropna(subset=['Month']).copy()
            if working_df.empty:
                return pd.DataFrame()

            monthly_agg = working_df.groupby('Month').agg(
                Impressions=('_impressions', 'sum'),
                ATC=('_atc', 'sum'),
                Orders=('_orders', 'sum'),
                Spends=('_budget_consumed', 'sum'),
                Sales=('_sales', 'sum')
            ).reset_index()

            monthly_agg['CPM'] = monthly_agg.apply(lambda r: round((r['Spends'] / r['Impressions']) * 1000) if r['Impressions'] > 0 else 0, axis=1)
            monthly_agg['ROAS'] = monthly_agg.apply(lambda r: round(r['Sales'] / r['Spends'], 2) if r['Spends'] > 0 else 0.0, axis=1)
            monthly_agg['ACOS'] = monthly_agg.apply(lambda r: round((r['Spends'] / r['Sales']) * 100, 2) if r['Sales'] > 0 else 0.0, axis=1)

            monthly_agg['Spends'] = monthly_agg['Spends'].round(2)
            monthly_agg['Sales'] = monthly_agg['Sales'].round(2)

            all_months = get_calendar_months(df_input)
            month_order = {m: i for i, m in enumerate(all_months)}
            monthly_agg['month_order'] = monthly_agg['Month'].map(
                lambda x: month_order.get(str(x).strip().upper(), 99)
            )
            monthly_agg = monthly_agg.sort_values('month_order').drop(columns=['month_order'])

            col_order = ['Month', 'Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']
            monthly_agg = monthly_agg[col_order].reset_index(drop=True)

            pct_rows = []
            if len(monthly_agg) >= 2:
                prev_row = monthly_agg.loc[len(monthly_agg) - 2]
                curr_row = monthly_agg.loc[len(monthly_agg) - 1]
                pct_row = {'Month': f"{curr_row['Month']} vs {prev_row['Month']} %"}
                for metric in ['Impressions', 'CPM', 'ATC', 'Orders', 'Spends', 'Sales', 'ROAS', 'ACOS']:
                    prev_val = prev_row[metric]
                    curr_val = curr_row[metric]
                    if prev_val > 0:
                        pct_change = round(((curr_val - prev_val) / prev_val) * 100)
                        pct_row[metric] = f"{pct_change}%" if pct_change <= 0 else f"+{pct_change}%"
                    else:
                        pct_row[metric] = "0%"
                pct_rows.append(pct_row)

            if pct_rows:
                pct_df = pd.DataFrame(pct_rows)
                monthly_agg = pd.concat([monthly_agg, pct_df], ignore_index=True)

            monthly_agg['ACOS'] = monthly_agg['ACOS'].apply(lambda v: f"{v:.2f}%" if isinstance(v, (int, float)) else str(v))
            return monthly_agg

        def _format_dashboard_value(value, column_name=""):
            if pd.isna(value): return value
            col = str(column_name).strip().upper()
            is_percentage = ('%' in col) or ('ACOS' in col)
            if isinstance(value, (int, np.integer)) and not isinstance(value, bool):
                return f"{int(value):,}" if not is_percentage else f"{int(value)}%"
            if isinstance(value, (float, np.floating)) and np.isfinite(value):
                if is_percentage: return f"{value:.2f}%"
                if float(value).is_integer(): return f"{int(value):,}"
                return f"{value:,.2f}"
            return value

        def format_dashboard_dataframe(df):
            out = df.copy()
            if isinstance(out.columns, pd.MultiIndex):
                for col in out.columns:
                    out[col] = out[col].map(lambda v: _format_dashboard_value(v, f"{col[0]} {col[1]}"))
            else:
                for col in out.columns:
                    out[col] = out[col].map(lambda v: _format_dashboard_value(v, col))
            return out

        def render_unified_single_table(df_to_show, key_prefix="mom"):
            if df_to_show is None or df_to_show.empty: return
            working_df = df_to_show.copy()
            bottom_rows = None
            if isinstance(working_df.index, pd.Index) and 'Grand Total' in working_df.index:
                main_df = working_df.drop('Grand Total')
                bottom_rows = working_df.loc[['Grand Total']]
            elif 'Month' in working_df.columns and any(('VS' in str(m).upper()) for m in working_df['Month'].values):
                pct_mask = working_df['Month'].apply(lambda x: 'VS' in str(x).upper())
                main_df = working_df[~pct_mask].copy()
                bottom_rows = working_df[pct_mask].copy()
            else:
                main_df = working_df
                bottom_rows = None

            sort_cols = [str(c) for c in main_df.columns] if not isinstance(main_df.columns, pd.MultiIndex) else [f"{c[0]} - {c[1]}" for c in main_df.columns]
            c_sort1, c_sort2 = st.columns([3, 1])
            with c_sort1:
                selected_sort_col = st.selectbox("Sort Table Column:", ["Default Order"] + sort_cols, key=f"{key_prefix}_sort_col")
            with c_sort2:
                sort_order = st.radio("Order:", ["Ascending", "Descending"], key=f"{key_prefix}_sort_ord", horizontal=True)

            if selected_sort_col != "Default Order":
                asc = (sort_order == "Ascending")
                if isinstance(main_df.columns, pd.MultiIndex):
                    idx_match = sort_cols.index(selected_sort_col)
                    col_key = main_df.columns[idx_match]
                    main_df = main_df.sort_values(by=col_key, ascending=asc)
                else:
                    main_df = main_df.sort_values(by=selected_sort_col, ascending=asc)

            if bottom_rows is not None and not bottom_rows.empty:
                unified_df = pd.concat([main_df, bottom_rows])
            else:
                unified_df = main_df

            def highlight_percentage_cells(val):
                val_str = str(val).strip()
                if '%' in val_str:
                    if val_str.startswith('-'): return 'background-color: #f8d7da; color: #721c24; font-weight: bold;'
                    elif val_str.startswith('+'): return 'background-color: #d4edda; color: #155724; font-weight: bold;'
                return ''

            if hasattr(unified_df.style, "map"):
                styled_df = unified_df.style.map(highlight_percentage_cells)
            else:
                styled_df = unified_df.style.applymap(highlight_percentage_cells)

            display_df = format_dashboard_dataframe(unified_df)
            st.dataframe(
                display_df.style.map(highlight_percentage_cells) if hasattr(display_df.style, "map") else display_df.style.applymap(highlight_percentage_cells),
                use_container_width=True,
                hide_index=False if isinstance(unified_df.index, pd.MultiIndex) or unified_df.index.name else True,
                key=f"{key_prefix}_single_unified_grid"
            )

        st.markdown("### 🔍 Global Dashboard Filters")
        available_months = ["All Months"] + get_calendar_months(final_df)
        selected_month = st.selectbox("Select Month Across Dashboard (Excluding Comparison Tables)", available_months)

        filtered_df = final_df.copy()
        if selected_month != "All Months":
            filtered_df = filtered_df[filtered_df['Month'] == selected_month]

        total_impressions = filtered_df['_impressions'].sum()
        total_sales = round(filtered_df['_sales'].sum(), 2)
        total_orders = filtered_df['_orders'].sum()
        total_atc = filtered_df['_atc'].sum()
        total_budget = round(filtered_df['_budget_consumed'].sum(), 2)
        overall_roas = round((total_sales / total_budget), 2) if total_budget > 0 else 0.0

        st.markdown("### 📈 Overall Swiggy Campaign Performance Dashboard")
        
        row1_col1, row1_col2, row1_col3 = st.columns(3)
        with row1_col1: st.metric("Total Impressions", f"{int(total_impressions):,}")
        with row1_col2: st.metric("Total GMV / Sales", f"₹{total_sales:,.2f}")
        with row1_col3: st.metric("Total Budget Burnt", f"₹{total_budget:,.2f}")

        row2_col1, row2_col2, row2_col3 = st.columns(3)
        with row2_col1: st.metric("Overall RoAS", f"{overall_roas:.2f}x")
        with row2_col2: st.metric("Total Conversions / Orders", f"{int(total_orders):,}")
        with row2_col3: st.metric("Total Add To Cart", f"{int(total_atc):,}")

        st.divider()

        st.markdown("### 📑 Navigation & Swiggy Analytics Tabs")
        main_tab1, main_tab2, main_tab3, main_tab4, main_tab5, main_tab6, main_tab7, main_tab8, main_tab9, main_tab10 = st.tabs([
            "📄 Raw Files Preview",
            "📌 Consolidated Master",
            "📊 Comparison Tables (MoM)",
            "📈 Interactive Trend Analytics",
            "🎯 Campaign Performance",
            "🏙️ City Performance",
            "📦 Product Performance",
            "📢 Ad Property / Placement",
            "🔎 Keyword Performance",
            "📅 Weekly Trend"
        ])

        with main_tab1:
            st.caption("Inspect individual sheets tab-by-tab for each uploaded file.")
            selected_file_name = st.selectbox("Select Uploaded File to Preview:", list(raw_files_dict.keys()))
            if selected_file_name:
                sheets = raw_files_dict[selected_file_name]
                selected_sheet = st.selectbox("Select Sheet Tab:", list(sheets.keys()))
                if selected_sheet:
                    st.dataframe(format_dashboard_dataframe(sheets[selected_sheet].head(100)), use_container_width=True, hide_index=True)

        with main_tab2:
            st.caption("Preview the combined dataset across all uploaded files.")
            st.dataframe(format_dashboard_dataframe(final_df.head(100)), use_container_width=True, hide_index=True)

        with main_tab3:
            st.subheader("📊 Month-on-Month Comparison Tables")
            monthly_summary_df = create_monthly_summary_table(final_df)
            camp_pivot = create_mom_comparison_table(final_df, 'CAMPAIGN_NAME') if 'CAMPAIGN_NAME' in final_df.columns else None
            city_pivot = create_mom_comparison_table(final_df, 'CITY') if 'CITY' in final_df.columns else None
            prod_pivot = create_mom_comparison_table(final_df, 'PRODUCT_NAME') if 'PRODUCT_NAME' in final_df.columns else None
            prop_pivot = create_mom_comparison_table(final_df, 'AD_PROPERTY') if 'AD_PROPERTY' in final_df.columns else None

            mom_dict = {
                "Monthly_Summary": monthly_summary_df,
                "Campaign_MoM": camp_pivot,
                "City_MoM": city_pivot,
                "Product_MoM": prod_pivot,
                "AdProperty_MoM": prop_pivot
            }
            all_pivots_bytes = convert_all_pivots_to_excel(mom_dict)

            st.download_button(
                label="📥 Download All MoM Comparison Tables (.xlsx)",
                data=all_pivots_bytes,
                file_name="Swiggy_All_MoM_Comparison_Tables.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key="btn_dl_all_mom"
            )

            st.divider()

            comp_sub_tab0, comp_sub_tab1, comp_sub_tab2, comp_sub_tab3, comp_sub_tab4 = st.tabs([
                "📅 Monthly Summary",
                "🎯 Campaign Comparison",
                "🏙️ City Comparison",
                "📦 Product Comparison",
                "📢 Ad Property Comparison"
            ])

            with comp_sub_tab0:
                if not monthly_summary_df.empty:
                    render_unified_single_table(monthly_summary_df, key_prefix="m_sum")
            with comp_sub_tab1:
                if camp_pivot is not None and not camp_pivot.empty:
                    render_unified_single_table(camp_pivot, key_prefix="c_pivot")
            with comp_sub_tab2:
                if city_pivot is not None and not city_pivot.empty:
                    render_unified_single_table(city_pivot, key_prefix="city_pivot")
            with comp_sub_tab3:
                if prod_pivot is not None and not prod_pivot.empty:
                    render_unified_single_table(prod_pivot, key_prefix="prod_pivot")
            with comp_sub_tab4:
                if prop_pivot is not None and not prop_pivot.empty:
                    render_unified_single_table(prop_pivot, key_prefix="prop_pivot")

        with main_tab4:
            st.subheader("📈 Interactive Multi-Metric Trend Analytics")
            all_df_months = get_calendar_months(final_df)
            selected_trend_months = st.multiselect("Select Months:", options=all_df_months, default=all_df_months, key="trend_m")

            metric_map = {
                'GMV / Sales (₹)': '_sales',
                'Budget Burnt (₹)': '_budget_consumed',
                'ROAS': 'ROAS',
                'Orders / Conversions': '_orders',
                'Add To Cart (ATC)': '_atc',
                'Impressions': '_impressions',
                'ACOS (%)': 'ACOS',
                'CPM (₹)': 'CPM'
            }

            selected_trend_metrics = st.multiselect("Select Metrics:", options=list(metric_map.keys()), default=['GMV / Sales (₹)', 'Budget Burnt (₹)', 'ROAS'], key="trend_met")

            if selected_trend_months and selected_trend_metrics:
                trend_df = final_df[final_df['Month'].isin(selected_trend_months)].copy()
                monthly_summary = trend_df.groupby('Month').agg(
                    _impressions=('_impressions', 'sum'),
                    _atc=('_atc', 'sum'),
                    _orders=('_orders', 'sum'),
                    _budget_consumed=('_budget_consumed', 'sum'),
                    _sales=('_sales', 'sum')
                ).reset_index()

                monthly_summary['CPM'] = monthly_summary.apply(lambda r: round((r['_budget_consumed'] / r['_impressions']) * 1000) if r['_impressions'] > 0 else 0, axis=1)
                monthly_summary['ROAS'] = monthly_summary.apply(lambda r: round(r['_sales'] / r['_budget_consumed'], 2) if r['_budget_consumed'] > 0 else 0.0, axis=1)
                monthly_summary['ACOS'] = monthly_summary.apply(lambda r: round((r['_budget_consumed'] / r['_sales']) * 100, 2) if r['_sales'] > 0 else 0.0, axis=1)

                month_order = {m: i for i, m in enumerate(all_df_months)}
                monthly_summary['month_order'] = monthly_summary['Month'].map(lambda x: month_order.get(str(x).strip().upper(), 99))
                monthly_summary = monthly_summary.sort_values('month_order').drop(columns=['month_order'])

                fig_trend = make_subplots(specs=[[{"secondary_y": True}]])
                metric_colors = {'GMV / Sales (₹)': '#1E88E5', 'Budget Burnt (₹)': '#90CAF9', 'ROAS': '#1565C0', 'Orders / Conversions': '#26A69A', 'Add To Cart (ATC)': '#7986CB', 'Impressions': '#4DB6AC', 'ACOS (%)': '#AB47BC', 'CPM (₹)': '#8D6E63'}
                bar_metrics = {'GMV / Sales (₹)', 'Budget Burnt (₹)'}

                for metric_label in selected_trend_metrics:
                    col_key = metric_map[metric_label]
                    use_sec_y = metric_label in ['ROAS', 'ACOS (%)', 'CPM (₹)']
                    color = metric_colors.get(metric_label, '#1E88E5')
                    values = monthly_summary[col_key]

                    if metric_label in bar_metrics:
                        fig_trend.add_trace(go.Bar(x=monthly_summary['Month'], y=values, name=metric_label, marker=dict(color=color)), secondary_y=use_sec_y)
                    else:
                        fig_trend.add_trace(go.Scatter(x=monthly_summary['Month'], y=values, name=metric_label, mode='lines+markers', line=dict(shape='spline', width=3.5, color=color)), secondary_y=use_sec_y)

                fig_trend.update_layout(title="<b>Swiggy Multi-Metric Trend Analysis</b>", template='plotly_white', height=550)
                st.plotly_chart(fig_trend, use_container_width=True)

        with main_tab5:
            if 'CAMPAIGN_NAME' in filtered_df.columns:
                c_opts = ["All"] + sorted([str(x) for x in filtered_df['CAMPAIGN_NAME'].dropna().unique()])
                sel_c = st.selectbox("Select Campaign Name:", c_opts)
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, 'CAMPAIGN_NAME', sel_c)), use_container_width=True, hide_index=True)

        with main_tab6:
            if 'CITY' in filtered_df.columns:
                city_opts = ["All"] + sorted([str(x) for x in filtered_df['CITY'].dropna().unique()])
                sel_city = st.selectbox("Select City:", city_opts)
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, 'CITY', sel_city)), use_container_width=True, hide_index=True)

        with main_tab7:
            if 'PRODUCT_NAME' in filtered_df.columns:
                prod_opts = ["All"] + sorted([str(x) for x in filtered_df['PRODUCT_NAME'].dropna().unique()])
                sel_prod = st.selectbox("Select Product Name:", prod_opts)
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, 'PRODUCT_NAME', sel_prod)), use_container_width=True, hide_index=True)

        with main_tab8:
            if 'AD_PROPERTY' in filtered_df.columns:
                prop_opts = ["All"] + sorted([str(x) for x in filtered_df['AD_PROPERTY'].dropna().unique()])
                sel_prop = st.selectbox("Select Ad Property:", prop_opts)
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, 'AD_PROPERTY', sel_prop)), use_container_width=True, hide_index=True)

        with main_tab9:
            kw_col = 'KEYWORD' if 'KEYWORD' in filtered_df.columns else ('PRODUCT_NAME' if 'PRODUCT_NAME' in filtered_df.columns else None)
            if kw_col:
                kw_opts = ["All"] + sorted([str(x) for x in filtered_df[kw_col].dropna().unique()])
                sel_kw = st.selectbox(f"Select {kw_col}:", kw_opts)
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, kw_col, sel_kw)), use_container_width=True, hide_index=True)

        with main_tab10:
            if 'Week' in filtered_df.columns and filtered_df['Week'].notna().any():
                st.dataframe(format_dashboard_dataframe(compute_grouped_table(filtered_df, 'Week', "All")), use_container_width=True, hide_index=True)

        st.divider()
        st.subheader("💾 Download Consolidated Master Workbook")
        buffer_multi = io.BytesIO()
        with pd.ExcelWriter(buffer_multi, engine='openpyxl') as writer:
            for raw_tab_name, df_list in consolidated_raw_tabs.items():
                combined_raw_tab_df = pd.concat(df_list, ignore_index=True)
                combined_raw_tab_df.to_excel(writer, sheet_name=raw_tab_name[:31], index=False)
            final_df.to_excel(writer, sheet_name='Consolidated_Master', index=False)
        buffer_multi.seek(0)
        st.download_button(
            label="📥 Download Complete Excel Workbook (.xlsx)",
            data=buffer_multi,
            file_name="Swiggy_Consolidated_Master_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
